"""
Teleprompter Script Generator: SOP file -> {stem}.teleprompter.md

The output file contains the full teleprompter script followed by an SME review
checklist separated by a horizontal rule.

Usage:
    python teleprompter.py --sop inbox/en/my-procedure.sop.md \
                           --output results/en/my-procedure \
                           --ollama-host localhost:11434 \
                           --model qwen2.5:7b
"""

import argparse
import json
import pathlib
import re
import sys

from ollama_utils import _call_ollama


NUMBERS = [
    "ONE", "TWO", "THREE", "FOUR", "FIVE", "SIX", "SEVEN", "EIGHT", "NINE", "TEN",
    "ELEVEN", "TWELVE", "THIRTEEN", "FOURTEEN", "FIFTEEN", "SIXTEEN", "SEVENTEEN",
    "EIGHTEEN", "NINETEEN", "TWENTY",
]

REVIEW_CHECKLIST = """\
# SME Review Checklist

Complete this checklist before recording begins.

- [ ] Presenter name filled in (search for `[Presenter Name]` in the teleprompter script)
- [ ] All `[bracketed placeholders]` resolved
- [ ] Step sequence matches the current version of the written S-O-P
- [ ] Warnings and compliance notes included where required
- [ ] Estimated runtime is acceptable for the intended audience
- [ ] Visual cues are achievable with available equipment or software
- [ ] Sign-off: Reviewed by: _______________________ Date: _______________
"""


# ---------------------------------------------------------------------------
# Input parsing
# ---------------------------------------------------------------------------

def _read_sop(path: pathlib.Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            print("Error: pypdf is required to read PDF files. pip install pypdf", file=sys.stderr)
            sys.exit(1)
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        try:
            import docx
        except ImportError:
            print("Error: python-docx is required to read .docx files. pip install python-docx", file=sys.stderr)
            sys.exit(1)
        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    # fallback: plain text
    return path.read_text()


# ---------------------------------------------------------------------------
# Prompt
# ---------------------------------------------------------------------------

def build_prompt(sop_text: str) -> str:
    return (
        "You are an expert video script writer specialising in teleprompter scripts.\n"
        "Convert the Standard Operating Procedure below into a structured video script.\n\n"
        "FORMATTING RULES — follow these exactly:\n"
        "- Write for the ear, not the eye. Use conversational English.\n"
        "- Keep all sentences under 15 words.\n"
        "- Write out numbers in full (e.g. 'one thousand dollars', not '$1,000').\n"
        "- Write out symbols in full (e.g. 'ten percent', not '10%').\n"
        "- Write out dates in full (e.g. 'September 27th').\n"
        "- Hyphenate acronyms spoken as letters (e.g. 'S-O-P', 'C-R-M', 'I-D').\n"
        "- Do NOT use bullet points or bold text inside audio fields.\n"
        "- Use \\n line breaks at natural breath points — not at sentence ends, but mid-thought where a presenter would pause.\n"
        "- Add [PAUSE] on its own line where the presenter needs a deliberate beat between ideas.\n\n"
        "SOP TEXT:\n"
        f"{sop_text}\n\n"
        "Return a JSON object with exactly these keys:\n"
        '- "sop_title": title of the SOP (string)\n'
        '- "video_goal": one sentence — what the viewer will achieve (string)\n'
        '- "prerequisites": list of 1-3 items needed before starting (list of strings)\n'
        '- "hook_audio": the spoken content that follows immediately after "Hi, I\'m [Presenter Name]." — CRITICAL: do NOT begin with any greeting, name, or self-introduction of any kind — no "Hi", "Hello", "Hey", "I\'m [name]", "My name is", or any variation; the presenter introduction is already written for you and will be prepended automatically; begin the very first word with an engaging question or relatable statement for the audience, then organically weave in what they will achieve and what they need beforehand; use breath-point line breaks and [PAUSE] markers (string)\n'
        '- "steps": list of {"name": string, "visual_cue": string, "audio": string} — one entry per SOP step; each audio should be 2-4 sentences with breath-point line breaks (list)\n'
        '- "has_pitfalls": true if the SOP contains warnings, notes, or troubleshooting (boolean)\n'
        '- "pitfalls_visual_cue": visual cue description for the pitfalls section (string)\n'
        '- "pitfalls_audio": teleprompter text for pitfalls, empty string if has_pitfalls is false (string)\n'
        '- "outro_audio": full closing in natural spoken English — brief summary and encouragement; do NOT include the helpdesk/contact line — that is added automatically (string)\n'
        '- "department_or_helpdesk": team or contact for help, inferred from the SOP (string)\n'
    )


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------

def _strip_formatting(text: str) -> str:
    """Remove bold markers and leading bullet points from audio text.
    Ensures [PAUSE] markers are surrounded by blank lines."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"^[\-\*]\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"\s*(\[PAUSE\])\s*", r"\n\n\1\n\n", text)
    return text.strip()


def _strip_greeting(text: str) -> str:
    """Strip any leading greeting sentence the LLM adds to hook_audio."""
    return re.sub(
        r"^(Hi|Hello|Hey),?\s+I'?m\s+\S+[.!]?\s*",
        "",
        text,
        flags=re.IGNORECASE,
    ).strip()


def _count_audio_words(data: dict) -> int:
    parts = [
        data.get("hook_audio", ""),
        data.get("pitfalls_audio", ""),
        data.get("outro_audio", ""),
    ]
    for step in data.get("steps", []):
        parts.append(step.get("audio", ""))
    return sum(len(p.split()) for p in parts if p)


def _runtime_estimate(word_count: int) -> str:
    seconds = round((word_count / 140) * 60 / 30) * 30
    minutes = seconds // 60
    secs = seconds % 60
    if minutes == 0:
        return f"~{secs} sec"
    if secs == 0:
        return f"~{minutes} min"
    return f"~{minutes} min {secs} sec"


def _step_number(i: int) -> str:
    if i <= len(NUMBERS):
        return NUMBERS[i - 1]
    return str(i)


# ---------------------------------------------------------------------------
# Markdown renderer
# ---------------------------------------------------------------------------

def _render_markdown(data: dict) -> str:
    sop_title = data.get("sop_title", "[SOP Title]")
    video_goal = data.get("video_goal", "[Video Goal]")
    hook_audio = _strip_greeting(_strip_formatting(data.get("hook_audio", "")))
    steps = data.get("steps", [])
    has_pitfalls = data.get("has_pitfalls", False)
    pitfalls_visual_cue = data.get("pitfalls_visual_cue", "Warning Icon on Screen")
    pitfalls_audio = _strip_formatting(data.get("pitfalls_audio", ""))
    outro_audio = _strip_formatting(data.get("outro_audio", ""))
    helpdesk = data.get("department_or_helpdesk", "[Department/Helpdesk]")

    word_count = _count_audio_words(data)
    runtime = _runtime_estimate(word_count)

    outro_section = 4 if (has_pitfalls and pitfalls_audio) else 3

    lines = [
        f"# {sop_title} — Teleprompter Video Script",
        "",
        f"**SOP Title:** {sop_title}  ",
        f"**Video Goal:** {video_goal}  ",
        f"**Estimated Runtime:** {runtime} (at 130–150 words per minute)  ",
        "**Target Tone:** Professional, encouraging, and clear",
        "",
        "---",
        "",
        "## SECTION 1 — Hook & Intro",
        "",
        f'VISUAL: Title card — "{sop_title}" / Presenter on camera',
        "",
        "---",
        "",
        "Hi, I'm [Presenter Name].",
        "",
        hook_audio,
        "",
        "Let's get into it.",
        "",
        "---",
        "",
        "## SECTION 2 — Step-by-Step Walkthrough",
        "",
        "---",
        "",
    ]

    for i, step in enumerate(steps, start=1):
        num = _step_number(i)
        name = step.get("name", f"Step {i}")
        visual_cue = step.get("visual_cue", "Screen recording")
        audio = _strip_formatting(step.get("audio", ""))
        lines += [
            f"### STEP {num} — {name}",
            "",
            f"VISUAL: {visual_cue}",
            "",
            "---",
            "",
            audio,
            "",
            "---",
            "",
        ]

    if has_pitfalls and pitfalls_audio:
        lines += [
            "## SECTION 3 — Common Pitfalls & Troubleshooting",
            "",
            f"VISUAL: {pitfalls_visual_cue}",
            "",
            "---",
            "",
            pitfalls_audio,
            "",
            "---",
            "",
        ]

    lines += [
        f"## SECTION {outro_section} — Outro & Call to Action",
        "",
        "VISUAL: Presenter back on camera",
        "",
        "---",
        "",
        outro_audio,
        "",
        "For more details, you can read the full written S-O-P linked below this video.",
        f"If you run into any blockers, reach out to the {helpdesk} team.",
        "",
        "Thanks for watching.",
        "",
        "---",
        "",
        REVIEW_CHECKLIST.strip(),
        "",
    ]

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def generate(
    sop_text: str,
    output_stem: pathlib.Path,
    ollama_host: str = "localhost:11434",
    model: str = "qwen2.5:7b",
) -> pathlib.Path:
    """Generate teleprompter script from SOP text.

    Writes a single file: {output_stem}.teleprompter.md
    The file contains the full script followed by the SME review checklist.
    Returns script_path.
    """
    script_path = output_stem.parent / (output_stem.name + ".teleprompter.md")

    if script_path.exists():
        return script_path

    prompt = build_prompt(sop_text)
    raw = _call_ollama(prompt, model, ollama_host)

    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        script_path.write_text(f"<!-- WARNING: raw LLM output -->\n\n{raw}\n")
        return script_path

    md_text = _render_markdown(data)
    script_path.write_text(md_text)

    return script_path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Generate a teleprompter script from an SOP.")
    parser.add_argument("--sop", required=True, help="Path to SOP file (.md, .txt, .pdf, .docx)")
    parser.add_argument("--output", required=True, help="Output stem (e.g. results/en/myvideo)")
    parser.add_argument("--ollama-host", default="localhost:11434")
    parser.add_argument("--model", default="qwen2.5:7b")
    args = parser.parse_args()

    sop_path = pathlib.Path(args.sop)
    if not sop_path.exists():
        print(f"Error: SOP file not found: {sop_path}", file=sys.stderr)
        sys.exit(1)

    sop_text = _read_sop(sop_path)
    output_stem = pathlib.Path(args.output)
    output_stem.parent.mkdir(parents=True, exist_ok=True)

    script_path = generate(
        sop_text=sop_text,
        output_stem=output_stem,
        ollama_host=args.ollama_host,
        model=args.model,
    )
    print(f"Script  -> {script_path}")


if __name__ == "__main__":
    main()
